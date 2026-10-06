"""Tạo file Excel (.xlsx) để upload lên Google Sheets từ các CSV trong resources/.

Chạy:  python3 scripts/tao-sheet-de-thi.py        (cần: pip install openpyxl)
Vào:   resources/phanthi/phan1..5.csv, resources/info.csv, resources/ds_hoc_vien.csv
Ra:    resources/de-thi-2026.xlsx với các tab CauHoi, PhanThi, HocVien, ThongTin, KetQua (định dạng như README).
"""
import csv
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

GOC = Path(__file__).resolve().parent.parent / 'resources'
RA = GOC / 'de-thi-2026.xlsx'

# Mỗi file phan{n}.csv ↔ một nhóm câu hỏi, tiền tố mã câu, loại câu hỏi hệ thống.
# Phần V: Câu 1–3 điền vào chỗ trống (DK), Câu 4 tự luận dài (TL); phần này không xáo thứ tự.
PHAN = {
    1: dict(nhom='giao_ly_hon_nhan', tien_to='HN', loai='TN'),
    2: dict(nhom='giao_ly_can_ban', tien_to='CB', loai='TN'),
    3: dict(nhom='kinh_bon', tien_to='KB', loai='TN'),
    4: dict(nhom='dien_khuyet', tien_to='DK', loai='DK'),
    5: dict(nhom='song_hon_nhan', tien_to='SH', loai='DK', xao_tron=False),
}
TEN_THANH = ['Giuse Maria', 'Vinh Sơn', 'Đa Minh', 'Gieronimo', 'Phaolô', 'Giuse', 'Maria', 'Anna']  # dài trước ngắn sau


def loi(msg):
    sys.exit(f'LỖI: {msg}')


def doc_csv(path):
    with open(path, encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


# ---- info.csv: tiêu đề phần, hướng dẫn, số câu/điểm; thông tin đề thi
info = doc_csv(GOC / 'info.csv')
phan_info = {}  # tiêu đề phần -> {'huong_dan': str, 'so_cau': int|None, 'diem': int|None}
thong_tin = []  # (khóa, nội dung)
dong_dau = []
da_thay_tieu_de = da_thay_luu_y = False
for r in info:
    muc, nd = r['Mục'].strip(), r['Nội dung'].strip()
    if not muc and not nd:
        continue
    if muc.startswith('PHẦN'):
        p = phan_info.setdefault(muc, dict(huong_dan='', so_cau=None, diem=None))
        m = re.fullmatch(r'\((\d+) câu\s*[–-]\s*(\d+) điểm\)', nd)
        if m:
            p['so_cau'], p['diem'] = int(m[1]), int(m[2])
        else:
            p['huong_dan'] = nd
        continue
    # "Thông tin đề thi"
    if nd.startswith('ĐỀ THI'):
        thong_tin += [('to_chuc', x) for x in dong_dau] + [('tieu_de', nd)]
        da_thay_tieu_de = True
    elif (m := re.search(r'thời gian.*?(\d+)\s*phút', nd, re.I)):
        thong_tin.append(('thoi_luong_phut', m[1]))
    elif nd.rstrip(':').strip().lower() == 'lưu ý':
        da_thay_luu_y = True
    elif da_thay_luu_y:
        thong_tin.append(('luu_y', nd))
    elif not da_thay_tieu_de:
        dong_dau.append(nd)
if not da_thay_tieu_de or not any(k == 'thoi_luong_phut' for k, _ in thong_tin):
    loi('info.csv thiếu dòng "ĐỀ THI ..." hoặc "thời gian làm bài ... phút".')

# ---- câu hỏi
cau_hoi, phan_thi = [], []
tong_diem = 0
for n, cfg in PHAN.items():
    rows = doc_csv(GOC / 'phanthi' / f'phan{n}.csv')
    tieu_de = rows[0]['Phần'].strip()
    pi = phan_info.get(tieu_de)
    if not pi:
        loi(f'phan{n}.csv: tiêu đề "{tieu_de}" không có trong info.csv.')
    diem_moi = pi['diem'] / pi['so_cau'] if pi['so_cau'] else None
    if pi['so_cau'] and pi['so_cau'] != len(rows):
        loi(f'phan{n}.csv có {len(rows)} câu nhưng info.csv ghi {pi["so_cau"]} câu.')
    for i, r in enumerate(rows, 1):
        nd = r['Nội dung câu hỏi'].strip()
        diem = diem_moi
        # Phần V ghi điểm ngay trong câu: "(2 điểm)" hoặc "(viết khoảng 150-200 từ - 14 điểm)"
        if (m := re.search(r'\s*[-–]\s*(\d+) điểm\)\s*$', nd)):
            diem, nd = int(m[1]), nd[: m.start()] + ')'
        elif (m := re.search(r'\s*\((\d+) điểm\)\s*$', nd)):
            diem, nd = int(m[1]), nd[: m.start()]
        if diem is None:
            loi(f'phan{n}.csv dòng {i}: không xác định được điểm.')
        loai = cfg['loai']
        if n == 5 and i == len(rows):
            loai = 'TL'
        dap_an, lua_chon = '', []
        if loai == 'TN':
            lua_chon = [r[f'Lựa chọn {k}'].strip() for k in 'ABCD']
            dap_an = r['Đáp án'].strip().upper()
            if dap_an not in 'ABCD' or len(dap_an) != 1 or not all(lua_chon):
                loi(f'phan{n}.csv dòng {i}: thiếu lựa chọn hoặc đáp án sai ("{dap_an}").')
        diem = int(diem) if float(diem).is_integer() else diem
        tong_diem += diem
        cau_hoi.append([f"{cfg['tien_to']}{i:02d}", loai, cfg['nhom'], nd, dap_an, diem, *lua_chon])
    xao = '' if cfg.get('xao_tron', True) else 'không'
    phan_thi.append([cfg['nhom'], tieu_de, len(rows), pi['huong_dan'], xao])
if tong_diem != 100:
    loi(f'tổng điểm các câu là {tong_diem}, không phải 100.')

# ---- học viên: tách Tên Thánh (nếu có) khỏi họ tên
hoc_vien = []
for dong in (GOC / 'ds_hoc_vien.csv').read_text(encoding='utf-8-sig').splitlines()[1:]:
    ten = ' '.join(dong.split(',')[0].split())
    if not ten:
        continue
    tt = next((t for t in TEN_THANH if ten.startswith(t + ' ')), '')
    hoc_vien.append([tt, ten[len(tt):].strip(), '', ''])

# ---- ghi file
KET_QUA = ['Mã lượt', 'Tên Thánh', 'Họ và tên', 'Giáo xứ', 'Giáo phận', 'Giờ bắt đầu', 'Giờ nộp', 'Nộp trễ', 'Số lần rời tab',
           'Điểm TN', 'Bài tự luận', 'Điểm TL', 'Tổng điểm', 'IP', 'Vị trí', 'Trình duyệt', 'Bài đã nộp']
SHEETS = {
    'CauHoi': (['Mã', 'Loại', 'Nhóm', 'Nội dung', 'Đáp án', 'Điểm', 'Lựa chọn A', 'Lựa chọn B', 'Lựa chọn C', 'Lựa chọn D'], cau_hoi,
               [10, 8, 18, 70, 9, 8, 40, 40, 40, 40]),
    'PhanThi': (['Nhóm', 'Tiêu đề phần', 'Số câu', 'Hướng dẫn', 'Xáo'], phan_thi, [18, 48, 10, 32, 8]),
    'HocVien': (['Tên Thánh', 'Họ và tên', 'Giáo xứ', 'Giáo phận'], hoc_vien, [14, 32, 20, 20]),
    'ThongTin': (['Khóa', 'Nội dung'], thong_tin, [18, 90]),
    'KetQua': (KET_QUA, [], [38, 12, 28, 18, 18, 20, 20, 9, 10, 9, 60, 9, 11, 16, 24, 40, 90]),
}
wb = Workbook()
wb.remove(wb.active)
dam = Font(bold=True)
nen = PatternFill('solid', fgColor='E8ECF8')
for ten, (tieu_de, rows, rong) in SHEETS.items():
    ws = wb.create_sheet(ten)
    ws.append(tieu_de)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font, c.fill = dam, nen
    ws.freeze_panes = 'A2'
    for i, w in enumerate(rong, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
wb.save(RA)

dem = {t: sum(1 for r in cau_hoi if r[1] == t) for t in ('TN', 'DK', 'TL')}
print(f'Đã tạo {RA}')
print(f'  CauHoi: {len(cau_hoi)} câu {dem}, tổng {tong_diem} điểm | PhanThi: {len(phan_thi)} phần | HocVien: {len(hoc_vien)} '
      f'({sum(1 for h in hoc_vien if h[0])} có Tên Thánh) | ThongTin: {len(thong_tin)} dòng')
